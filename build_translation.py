#!/usr/bin/env python3
"""Build the English translation of `پیوست ۲.docx` (Appendix 2 – Final Research Proposal)
with page numbers in the footer. Output: `Appendix 2 - English.docx`."""

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
SRC = ROOT / "پیوست ۲.docx"
OUT = ROOT / "Appendix 2 - English.docx"
MEDIA = ROOT / ".build_media"

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def extract_media():
    MEDIA.mkdir(exist_ok=True)
    with zipfile.ZipFile(SRC) as z:
        for name in ("word/media/image1.png", "word/media/image2.jpeg"):
            (MEDIA / Path(name).name).write_bytes(z.read(name))


def add_field(run, instr):
    """Insert a Word field (e.g. PAGE / NUMPAGES) into a run."""
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        run._r.append(el)


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_page_border(section):
    sectPr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side, val in (("top", "thinThickSmallGap"), ("left", "thinThickSmallGap"),
                      ("bottom", "thickThinSmallGap"), ("right", "thickThinSmallGap")):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), val)
        b.set(qn("w:sz"), "24")
        b.set(qn("w:space"), "15")
        b.set(qn("w:color"), "auto")
        borders.append(b)
    sectPr.append(borders)


def set_table_borders(table, sz="4"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), sz)
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "auto")
        borders.append(b)
    tblPr.append(borders)


def para(doc, text="", style=None, bold=False, italic=False, size=None,
         align=None, space_after=6, first_indent=None, keep_next=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if keep_next:
        pf.keep_with_next = True
    return p


def labelled(doc, label, value, size=11, space_after=4):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(size)
    r2 = p.add_run(value)
    r2.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bibliography(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        # items are lists of (text, italic) tuples
        for text, ital in it:
            r = p.add_run(text)
            r.italic = ital
            r.font.size = Pt(11)


# ---------------------------------------------------------------------------
# document
# ---------------------------------------------------------------------------

def build():
    extract_media()
    doc = Document()

    # -- base styles -------------------------------------------------------
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in (("Heading 1", 14, "1F3864"), ("Heading 2", 12, "2E5395")):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
        # make sure east-asian/complex-script fonts follow too
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), "Calibri")

    # -- page setup (A4, same margins as the original) ---------------------
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.95)
    sec.footer_distance = Cm(1.1)
    set_page_border(sec)

    # -- footer with page numbers -----------------------------------------
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    r.font.size = Pt(10)
    r = fp.add_run()
    r.font.size = Pt(10)
    add_field(r, "PAGE")
    r = fp.add_run(" of ")
    r.font.size = Pt(10)
    r = fp.add_run()
    r.font.size = Pt(10)
    add_field(r, "NUMPAGES")

    # -- page 1: letterhead -----------------------------------------------
    head = doc.add_table(rows=1, cols=3)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    head.autofit = False
    widths = (Cm(6.0), Cm(5.4), Cm(6.0))
    for i, w in enumerate(widths):
        head.columns[i].width = w
        head.rows[0].cells[i].width = w

    # left: seminary info box
    c = head.rows[0].cells[0]
    c.paragraphs[0].text = ""
    box = c.add_table(rows=1, cols=1)
    set_table_borders(box, "8")
    bc = box.rows[0].cells[0]
    bp = bc.paragraphs[0]
    bp.paragraph_format.space_after = Pt(0)
    r = bp.add_run("Seminary: ")
    r.bold = True
    r.font.size = Pt(9.5)
    r = bp.add_run("Haj Mulla Sadeq Mojtahed Qomi")
    r.font.size = Pt(9.5)
    bp2 = bc.add_paragraph()
    bp2.paragraph_format.space_after = Pt(0)
    r = bp2.add_run("City: ")
    r.bold = True
    r.font.size = Pt(9.5)
    r = bp2.add_run("Qom          ")
    r.font.size = Pt(9.5)
    r = bp2.add_run("Province: ")
    r.bold = True
    r.font.size = Pt(9.5)
    r = bp2.add_run("Qom")
    r.font.size = Pt(9.5)

    # centre: "In the Name of God, the Exalted" (image from the original)
    c = head.rows[0].cells[1]
    cp = c.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run().add_picture(str(MEDIA / "image1.png"), width=Cm(1.5))
    cp2 = c.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp2.add_run("In the Name of God, the Exalted")
    r.italic = True
    r.font.size = Pt(9)

    # right: seminary management centre logo
    c = head.rows[0].cells[2]
    cp = c.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cp.add_run().add_picture(str(MEDIA / "image2.jpeg"), width=Cm(2.6))
    cp2 = c.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = cp2.add_run("Islamic Seminaries Management Centre\nDeputy of Research")
    r.font.size = Pt(8)

    para(doc, "", space_after=2)
    para(doc, "Final Research Proposal", bold=True, size=18,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, "(Appendix 2)", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=10)

    # -- student / project info box ---------------------------------------
    info = doc.add_table(rows=1, cols=1)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info, "12")
    ic = info.rows[0].cells[0]
    ic.width = Cm(17.4)
    ip = ic.paragraphs[0]
    ip.paragraph_format.space_after = Pt(4)
    r = ip.add_run("Full Name: ")
    r.bold = True
    ip.add_run("Mohammad Amin Sedghi")
    ip.add_run("\t\t\t\t")
    r = ip.add_run("File No.: ")
    r.bold = True
    ip.add_run("198162")
    p = ic.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Title of the Final Research: ")
    r.bold = True
    p.add_run("A Literary Analysis of ")
    r = p.add_run("Thānī Ithnayn")
    r.italic = True
    p.add_run(" (“the Second of Two”) and Similar Numerical Constructions")
    p = ic.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Supervisor: ")
    r.bold = True
    p.add_run("Hujjat al-Islam Mohammad Fazel Azimi")
    p = ic.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Date of submission of the detailed proposal: ")
    r.bold = True
    p.add_run("......... / ....... / ...............")
    p.add_run("\t\t")
    r = p.add_run("Signature: ")
    r.bold = True
    p.add_run("..................")

    para(doc, "", space_after=4)

    # -- 1. Statement of the problem --------------------------------------
    doc.add_heading("1. Statement of the Problem", level=1)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.add_run(
        "Among the special numerical constructions found in the Holy Qur’an and in the Arabic "
        "language that have rarely been analysed in detail in the textbooks of the Islamic "
        "seminaries are constructions such as "
    )
    p.add_run("thānī ithnayn").italic = True
    p.add_run(
        " (“the second of two”). From the morphological, syntactic and lexical points of view, "
        "these numerical constructions possess subtleties that play a decisive role in "
        "understanding their precise meaning and translating them correctly. Neglecting to "
        "examine the structure and function of these constructions can lead to ambiguity in "
        "exegesis or to errors in translation. It is therefore essential to present a "
        "comprehensive and well-documented analysis of such constructions so that seminary "
        "students and researchers, relying on the precise foundations of the Arabic language, "
        "can attain a more accurate and scholarly understanding of their syntactic and semantic "
        "context."
    )

    # -- 2. Necessity ------------------------------------------------------
    doc.add_heading("2. Necessity of the Research", level=1)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.add_run(
        "According to the survey carried out, although the numerical construction "
    )
    p.add_run("thānī ithnayn").italic = True
    p.add_run(
        " and similar expressions occur many times in religious texts, including the Holy "
        "Qur’an, no work or writing was found that offers a detailed and precise literary "
        "analysis of them. This research therefore undertakes a literary study and analysis of "
        "these constructions and can fill the gap and shortcomings that exist in this area of "
        "research."
    )

    # -- 3. Literature review ---------------------------------------------
    doc.add_heading("3. Review of the Literature", level=1)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.add_run(
        "In the survey that was carried out, no study, book, thesis or article was found that "
        "subjects the topic of "
    )
    p.add_run("thānī ithnayn").italic = True
    p.add_run(
        " to a literary examination; this can be counted among the advantages of the present "
        "research."
    )
    para(doc,
         "Nevertheless, some works have dealt with the number briefly in the course of other "
         "discussions. A report on these works follows:",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.add_heading("3.1. Morphological works", level=2)
    bibliography(doc, [
        [("Daqr, ‘Abd al-Ghanī, ", False), ("Mu‘jam al-Qawā‘id al-‘Arabiyya", True),
         (", Qom: al-Ḥamīd, 1410 AH, p. 293.", False)],
        [("Ḥasan, ‘Abbās, ", False),
         ("al-Naḥw al-Wāfī ma‘a Rabṭihi bi-l-Asālīb al-Rafī‘a wa-l-Ḥayāt al-Lughawiyya "
          "al-Mutajaddida", True),
         (", Tehran: Nāṣir Khusraw, 1367 SH, vol. 4, p. 515.", False)],
        [("Nāẓir al-Jaysh, Muḥammad ibn Yūsuf (together with Ibn Mālik, Muḥammad ibn "
          "‘Abd Allāh), ", False), ("Sharḥ al-Tashīl", True),
         (", ed. ‘Alī Muḥammad Fākhir and Jābir Muḥammad Barāja, Cairo: Dār al-Salām, "
          "1428 AH, vol. 5, p. 2457.", False)],
        [("Sībawayh, ‘Amr ibn ‘Uthmān (together with Yūsuf ibn Sulaymān al-A‘lam "
          "al-Shantamarī), ", False), ("al-Kitāb", True),
         (", Beirut: Mu’assasat al-A‘lamī li-l-Maṭbū‘āt, 1410 AH, vol. 2, 3rd printing.",
          False)],
    ])
    para(doc,
         "In the works listed above, the section on numbers refers only to the method of "
         "forming ordinal numbers.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.add_heading("3.2. Syntactic works", level=2)
    bibliography(doc, [
        [("Darwīsh, Muḥyī al-Dīn, ", False), ("I‘rāb al-Qur’ān al-Karīm wa-Bayānuhu", True),
         (", Damascus: Dār al-Yamāma, 1415 AH, vol. 4, p. 102, 4th printing.", False)],
        [("Ṣāfī, Maḥmūd, ", False),
         ("al-Jadwal fī I‘rāb al-Qur’ān wa-Ṣarfihi wa-Bayānihi ma‘a Fawā’id Naḥwiyya Hāmma",
          True),
         (", Damascus: Dār al-Rashīd, 1411 AH, vol. 10, p. 341.", False)],
        [("Ṣāfī, Maḥmūd, ", False),
         ("al-Jadwal fī I‘rāb al-Qur’ān wa-Ṣarfihi wa-Bayānihi ma‘a Fawā’id Naḥwiyya Hāmma",
          True),
         (", Damascus: Dār al-Rashīd, 1411 AH, vol. 6, p. 419.", False)],
    ])
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.add_run("The works listed above refer only to the grammatical parsing of the expressions ")
    p.add_run("thānī ithnayn").italic = True
    p.add_run(" under verse 40 of Sūrat al-Tawba and ")
    p.add_run("thālith thalātha").italic = True
    p.add_run(" (“the third of three”) under verse 73 of Sūrat al-Mā’ida.")

    doc.add_heading("3.3. Lexical works", level=2)
    bibliography(doc, [
        [("Qurashī, ‘Alī Akbar, ", False), ("Qāmūs-i Qur’ān", True),
         (", Tehran: Dār al-Kutub al-Islāmiyya, 1371 SH, vol. 1, p. 319.", False)],
        [("Azharī, Muḥammad ibn Aḥmad, ", False), ("Tahdhīb al-Lugha", True),
         (", introduction by Fāṭima Muḥammad Aṣlān, ed. ‘Umar Salāmī and ‘Abd al-Karīm Ḥāmid, "
          "supervised by Muḥammad ‘Awaḍ Mur‘ib, Beirut: Dār Iḥyā’ al-Turāth al-‘Arabī, "
          "1421 AH, vol. 15, p. 45.", False)],
        [("Ṭurayḥī, Fakhr al-Dīn ibn Muḥammad, ", False), ("Majma‘ al-Baḥrayn", True),
         (", ed. Aḥmad Ḥusaynī Ishkivarī, Tehran: Murtaḍavī Bookshop, 1375 SH, vol. 1, p. 73.",
          False)],
    ])
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.add_run("The works listed above refer only to the meanings of the expressions ")
    p.add_run("thānī ithnayn").italic = True
    p.add_run(" under verse 40 of Sūrat al-Tawba and ")
    p.add_run("thālith thalātha").italic = True
    p.add_run(" under verse 73 of Sūrat al-Mā’ida.")

    doc.add_heading("3.4. Exegetical works", level=2)
    bibliography(doc, [
        [("Ṭabāṭabā’ī, Muḥammad Ḥusayn, ", False), ("Tafsīr al-Mīzān", True),
         (" (Persian translation), trans. Muḥammad Bāqir Mūsavī, Qom: Society of Seminary "
          "Teachers of Qom, Islamic Publications Office, 1378 SH, vol. 9, p. 374.", False)],
    ])
    p = para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.add_run("This work, too, refers only to the meaning of the expression ")
    p.add_run("thānī ithnayn").italic = True
    p.add_run(" under verse 40 of Sūrat al-Tawba.")

    # -- 4. Hypothesis -----------------------------------------------------
    doc.add_heading("4. Research Hypothesis", level=1)
    para(doc,
         "It appears that, given the various morphological structures, some of the proposed "
         "meanings are correct and some are incorrect. Taking into account this variation in "
         "morphological structures and meanings, numerous combinations are obtained, some of "
         "which are valid and some invalid.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # -- 5. Questions ------------------------------------------------------
    doc.add_heading("5. Main and Subsidiary Research Questions", level=1)
    doc.add_heading("5.1. Main question", level=2)
    p = para(doc, style="List Bullet")
    p.add_run("What is the literary analysis of ")
    p.add_run("thānī ithnayn").italic = True
    p.add_run(" and similar numerical constructions?")

    doc.add_heading("5.2. Subsidiary questions", level=2)
    for kind in ("lexical", "morphological", "syntactic"):
        p = para(doc, style="List Bullet", space_after=2)
        p.add_run(f"What is the {kind} analysis of ")
        p.add_run("thānī ithnayn").italic = True
        p.add_run(" and similar numerical constructions?")

    # -- 6. Method ---------------------------------------------------------
    doc.add_heading("6. Research Method", level=1)
    para(doc,
         "This study employs the library (documentary) method; its data have been gathered from "
         "authoritative lexical, morphological and syntactic reference works. A report on the "
         "collected material is then presented in a descriptive–analytical manner: first the "
         "meanings of the words, then the morphological structure of the expressions in "
         "question, and finally their syntactic structure are examined, after which the "
         "analyses are tested for validity.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(doc,
         "This study is transmission-based (naqlī), in the sense that both the raw data and the "
         "analyses are drawn solely from the sources.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # -- 7. Organisation --------------------------------------------------
    doc.add_heading("7. Organisation of the Research (Chapters and Sub-chapters)", level=1)
    for item in ("Abstract", "Keywords"):
        p = para(doc, style="List Bullet", space_after=2)
        p.add_run(item).bold = True
    outline = [
        "Introduction",
        ("Conceptual clarification of the construction ", "thānī ithnayn", " and the like"),
        ("Semantic analysis of the construction ", "thānī ithnayn", " and the like"),
        ("Morphological structure of the construction ", "thānī ithnayn", " and the like"),
        ("Syntactic structure of the construction ", "thānī ithnayn", " and the like"),
        "Conclusion",
        "Bibliography",
    ]
    for item in outline:
        p = para(doc, style="List Number 2", space_after=2)
        if isinstance(item, tuple):
            p.add_run(item[0])
            p.add_run(item[1]).italic = True
            p.add_run(item[2])
        else:
            p.add_run(item)

    # -- 8. Sources table --------------------------------------------------
    doc.add_heading("8. Research Sources", level=1)
    headers = ["No.", "Title of Book / Article / Thesis", "Author", "Translator",
               "Place of Publication", "Publisher", "Year"]
    rows = [
        ["1", "Awḍaḥ al-Masālik ilā Alfiyyat Ibn Mālik",
         "Ibn Hishām, ‘Abd Allāh ibn Yūsuf", "", "Beirut", "al-Maktaba al-‘Aṣriyya", "1429 AH"],
        ["2", "Sharḥ Ibn ‘Aqīl", "Ibn ‘Aqīl, ‘Abd Allāh ibn ‘Abd al-Raḥmān", "",
         "[n.p.]", "[n.pub.]", ""],
        ["3", "Sharḥ al-Taṣrīḥ ‘alā al-Tawḍīḥ of Ibn Mālik on Syntax and Morphology, by "
               "Shaykh Abū Muḥammad ‘Abd Allāh ibn Yūsuf ibn Hishām al-Anṣārī",
         "Azharī, Khālid ibn ‘Abd Allāh", "", "Beirut", "Dār al-Fikr", ""],
        ["4", "al-Naḥw al-‘Arabī", "Barakāt, Ibrāhīm Ibrāhīm", "", "Cairo",
         "Dār al-Nashr li-l-Jāmi‘āt", "1428 AH"],
        ["5", "Kitāb al-Mudhakkar wa-l-Mu’annath", "Ibn al-Anbārī, Muḥammad ibn Qāsim", "",
         "Beirut", "Dār al-Rā’id al-‘Arabī", "1406 AH"],
    ]
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [Cm(1.1), Cm(5.2), Cm(3.6), Cm(1.9), Cm(2.0), Cm(2.6), Cm(1.6)]
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.width = col_w[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "D9E2F3")
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri].cells[ci]
            cell.width = col_w[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(val)
            r.font.size = Pt(9.5)
            if ci == 1 and val:
                r.italic = True
            if ci in (0, 6):
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # repeat header row
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)

    # -- Supervisor approval ----------------------------------------------
    para(doc, "", space_after=8)
    doc.add_heading("Supervisor’s Approval", level=1)
    para(doc,
         "I, the undersigned, .........................................., hereby approve the "
         "final research proposal of Mr. ........................................................",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(doc,
         "entitled: ...............................................................................................",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p = para(doc, space_after=12)
    p.add_run("Date:      /       /               ")
    p.add_run("\t\t\t\t")
    p.add_run("Signature: ..........................")

    # -- Department remarks -----------------------------------------------
    doc.add_heading("Supplementary Remarks of the Academic–Educational Department for "
                    "Improving the Research Proposal", level=1)
    para(doc, "." * 130 + "\n" + "." * 130 + "\n" + "." * 130 + "\n" + "." * 130 + "\n"
         + "." * 130 + "\n" + "." * 130, space_after=12)

    # -- Department decision ----------------------------------------------
    doc.add_heading("Opinion of the Academic–Educational Department on the Research Proposal",
                    level=1)
    para(doc,
         "The proposed research plan was examined at the meeting held on ............................ "
         "and the result is as follows:",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p = para(doc, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    for label in ("Approved", "Requires revision", "Rejected"):
        r = p.add_run("☐ ")
        r.font.size = Pt(13)
        r = p.add_run(label)
        r.bold = True
        p.add_run("          ")

    p = para(doc, space_after=8)
    r = p.add_run("Signatures of the members of the Academic–Educational Department:")
    r.bold = True
    p.add_run("        Date:        /          /          ")

    sig = doc.add_table(rows=4, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, title in enumerate(("Head of Department,", "Secretary of Department,")):
        c = sig.rows[0].cells[ci]
        c.text = ""
        r = c.paragraphs[0].add_run(title)
        r.bold = True
        for ri in range(1, 4):
            cell = sig.rows[ri].cells[ci]
            cell.text = ""
            cell.paragraphs[0].add_run("Full name / Signature: ....................................")
            cell.paragraphs[0].paragraph_format.space_after = Pt(10)

    doc.core_properties.title = "Final Research Proposal (Appendix 2) – English translation"
    doc.core_properties.author = "Mohammad Amin Sedghi"
    doc.core_properties.language = "en-GB"
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
